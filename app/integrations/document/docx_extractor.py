from __future__ import annotations

from pathlib import Path

from docx import Document

from app.integrations.document.base import (
    DocumentExtractor
)


class DocxExtractor(DocumentExtractor):

    def extrair(
        self,
        caminho: Path
    ) -> str:

        documento = Document(
            str(caminho)
        )

        partes: list[str] = []


        # =====================================================
        # PARÁGRAFOS
        # =====================================================

        for paragrafo in documento.paragraphs:

            texto = (
                paragrafo.text
                .strip()
            )

            if texto:

                partes.append(
                    texto
                )


        # =====================================================
        # TABELAS
        # =====================================================

        for tabela in documento.tables:

            for linha in tabela.rows:

                valores = []

                for celula in linha.cells:

                    texto = (
                        celula.text
                        .strip()
                    )

                    if texto:

                        valores.append(
                            texto
                        )

                if valores:

                    partes.append(
                        " | ".join(
                            valores
                        )
                    )


        texto_final = (
            "\n".join(
                partes
            )
            .strip()
        )


        if not texto_final:

            raise ValueError(
                "O arquivo DOCX não possui "
                "texto extraível."
            )


        return texto_final